import copy
import io
import re

import openpyxl
from PIL import Image

from ecommerce.consumption.excel import excel as X
from ecommerce.domain.product import Product
from ecommerce.ingestion.raw_store import RunStore, write_json
from ecommerce.platforms.shopee.parse import dataset as DS
from ecommerce.platforms.shopee.parse.listing import select_candidates
from helpers import PROFILE, make_cfg


def _jpeg() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (100, 100), (200, 60, 40)).save(buf, "JPEG")
    return buf.getvalue()


def _make_run(tmp_path, load):
    store = RunStore("shopee", "20260910", root=tmp_path)
    page = load("shopee_search_page.json")
    write_json(store.search_path(0), {"payload": page})
    kept, excluded = select_candidates([(0, page)], target=10, include_ads=False, profile=PROFILE)
    write_json(store.candidates_path, {"keyword": "giữ nhiệt", "sort_by": "sales",
                                       "kept": [c.dump() for c in kept],
                                       "excluded": [c.dump() for c in excluded]})
    raw = load("shopee_item_raw.json")
    write_json(store.item_path(11, 111), raw)
    # second product: search rank 3, almost empty pdp -> must still export
    raw2 = copy.deepcopy(raw)
    raw2["candidate"] = kept[1].dump()
    raw2["pdp"] = {"error": None, "data": {"item": {"item_id": 444, "shop_id": 44}}}
    raw2["ratings"] = []
    write_json(store.item_path(44, 444), raw2)
    # a poisoned file (blocked payload) must be reported, not exported
    write_json(store.item_path(55, 555), {"candidate": {"shopid": 55, "itemid": 555},
                                         "pdp": load("shopee_pdp_blocked.json")})
    store.record_failure("66_666", "ProductFailed: không bắt được pdp/get_pc trong 40s")
    return store


def test_export_end_to_end(tmp_path, load, monkeypatch):
    store = _make_run(tmp_path, load)
    monkeypatch.setattr(X, "fetch_many", lambda urls, *a, **k: {u: _jpeg() for u in urls})
    cfg = make_cfg(shopee={"target": 200}, export={"embedded_images": 3, "thumbnail_px": 110})
    dataset = DS.build_dataset(store, cfg)
    path = X.write_report(dataset, cfg.export, out_dir=tmp_path / "out", prefix=PROFILE.file_prefix)
    assert re.fullmatch(r"giu_nhiet_shopee_top2_20260910_\d{6}\.xlsx", path.name)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Shopee", "Detail", "Checklist đề bài", "Bị loại", "Lỗi crawl", "Thông tin"]
    ws = wb["Shopee"]
    headers = [c.value for c in ws[1]]
    # identity block: the first 3 columns stay visible when scrolling right (freeze_panes)
    assert headers[:7] == ["Hạng", "Ảnh 1", "Tên sản phẩm", "Link sản phẩm", "Mã SP",
                           "Loại sản phẩm", "Thương hiệu"]
    assert headers[-3:] == ["Link ảnh (tất cả)", "Ảnh 2", "Ảnh 3"]   # long cell + extra images go last
    for h in ("Chất liệu (trong / ngoài)", "Tỉ lệ phản hồi (%)", "Bán 30 ngày (TB tháng)",
              "Tồn kho", "Giá sau voucher", "Phân loại", "Link ảnh (tất cả)",
              "Ngày mở shop", "Link shop"):
        assert h in headers
    # dropped: outside the brief, near-constant, or misleading (featured 5-star reviews)
    for h in ("Hạng trên sàn", "Doanh thu ước tính 30 ngày", "Tuổi listing (tháng)", "Lượt thích",
              "Danh mục", "Chất liệu trong", "Chất liệu ngoài", "Link ảnh 1"):
        assert h not in headers
    assert not any(h.startswith("Đánh giá 1") for h in headers)
    assert not any(h.startswith("Phân loại 1 -") for h in headers)   # variants moved to own sheet
    assert len(headers) < 45
    col = {h: i + 1 for i, h in enumerate(headers)}
    assert ws.max_row == 3                                # 2 products, blocked one skipped
    assert [ws.cell(r, col["Hạng"]).value for r in (2, 3)] == [1, 2]
    assert ws.cell(2, col["Giá bán thấp nhất"]).value == 250000
    assert ws.cell(2, col["Giá sau voucher"]).value == 224000
    assert ws.cell(2, col["Chất liệu (trong / ngoài)"]).value == "Trong: Inox 316 · Ngoài: Inox 304"
    assert ws.cell(2, col["Link ảnh (tất cả)"]).value.count("\n") == 3      # 4 images, 1 cell
    assert ws.cell(2, col["Tồn kho"]).value is None                 # hidden by Shopee, not crawled
    assert ws.cell(2, col["Tình trạng"]).value.startswith("Còn hàng (2/3 SKU còn hàng)")
    assert ws.cell(2, col["Mã SP"]).value == "111"
    assert ws.cell(2, col["Link sản phẩm"]).hyperlink.target == "https://shopee.vn/product/11/111"
    assert len(ws._images) == 5                           # 3 + 2 (second product has 2 images)

    detail = wb["Detail"]
    vh = [c.value for c in detail[1]]
    # the user's target layout (screen recording 16:45)
    assert vh == ["Hạng", "Tên sản phẩm", "Ảnh", "Mã SP", "Mã SKU",
                  "Nhóm phân loại 1", "Giá trị 1", "Nhóm phân loại 2", "Giá trị 2",
                  "Dung tích", "Giá", "Giá gốc", "% giảm", "Giá sau voucher",
                  "Tồn kho (pieces available)", "Tình trạng", "Link Ảnh"]
    assert detail.max_row == 1 + 3
    row = [c.value for c in detail[3]]                   # 2nd SKU: "Đen" (black) / 750ml
    assert row[3:9] == ["111", "2", "Màu sắc", "Đen", "Dung tích", "750ml"]
    assert detail.cell(3, 17).hyperlink.target.endswith("/tvA")        # SKU's own picture
    assert len(detail._images) == 3                                     # one per SKU row

    checklist = {r[1].value: (r[3].value, r[4].value, r[5].value)
                 for r in wb["Checklist đề bài"].iter_rows(min_row=2)}
    assert len(checklist) == 19
    assert checklist["Giá"][1] == "✅ Đủ"
    assert checklist["Tồn kho"][1] == "⚠️ Một phần"       # status text only, no count
    assert "có tổng tồn kho đầy đủ" in checklist["Tồn kho"][2]

    failures = {r[0].value: r[1].value for r in wb["Lỗi crawl"].iter_rows(min_row=2)}
    assert set(failures) == {"55_555", "66_666"}
    excluded = [r[2].value for r in wb["Bị loại"].iter_rows(min_row=2)]
    assert any("Hộp cơm" in t for t in excluded)


def test_rank_orders_by_search_rank_and_caps():
    recs = [Product(platform="shopee", item_id=i, url="u", title="t", search_rank=r)
            for i, r in enumerate([5, 1, None, 3])]
    top = DS.rank(recs, target=3)
    assert [p.search_rank for p in top] == [1, 3, 5]
    assert [p.final_rank for p in top] == [1, 2, 3]


def test_control_characters_are_dropped_when_writing_excel(tmp_path, load, monkeypatch):
    """A shop-typed name with a control character -> Excel must show 'Trancy', not '_x0008_Trancy'.

    Seen for real when reconciling 200 Shopee products against the raw JSON (13/09):
    exactly 1 mismatched cell. Excel's XML forbids control characters, so
    xlsxwriter writes them as '_x0008_'.
    """
    from ecommerce.consumption.excel.excel import clean_text
    assert clean_text("Bình \x08Trancy\ttốt\ndòng 2") == "Bình Trancy\ttốt\ndòng 2"   # keeps tab and newline

    store = _make_run(tmp_path, load)
    raw = load("shopee_item_raw.json")
    raw["pdp"]["data"]["item"]["title"] = "Bình \x08Trancy giữ nhiệt"
    write_json(store.item_path(11, 111), raw)
    monkeypatch.setattr(X, "fetch_many", lambda urls, *a, **k: {})
    cfg = make_cfg()
    path = X.write_report(DS.build_dataset(store, cfg), cfg.export, out_dir=tmp_path / "out",
                          download_images=False)

    first_cells = [r[0] for r in openpyxl.load_workbook(path)["Shopee"].iter_rows(min_row=2, values_only=True)]
    ws = openpyxl.load_workbook(path)["Shopee"]
    title_col = [c.value for c in ws[1]].index("Tên sản phẩm") + 1
    titles = [ws.cell(r, title_col).value for r in range(2, ws.max_row + 1)]
    assert "Bình Trancy giữ nhiệt" in titles
    assert not any("_x0008_" in str(t) for t in titles)
    assert first_cells                # file is not empty


def test_every_excel_column_points_at_a_real_key(load):
    """A mistyped key in columns.py = an empty column that nobody reports.

    `excel.py` reads values with `r.get(key)`, so a wrong key just yields None: the
    column still shows, the header is still right, every cell is empty, and the whole
    suite stays green. Tried for real: renaming 'shop_location' to 'shop_locaton'
    -> 39/39 cells empty, 257 tests still passing.
    """
    from ecommerce.consumption.excel.columns import COLUMNS, SKU_COLUMNS
    from ecommerce.consumption.excel.excel import to_rows
    from ecommerce.platforms.shopee.parse.product import parse_product

    product = parse_product(load("shopee_item_raw.json"), PROFILE)
    row = to_rows([product])[0]
    missing = [c.key for c in COLUMNS if c.kind != "image" and c.key not in row]
    assert not missing, f"main sheet column points at a key missing from to_row(): {missing}"

    # two Detail-sheet keys are derived at write time and do not live on Variant
    DERIVED = {"discount_pct", "image_url"}                  # excel._sku_value / sku_image
    sku = product.variants[0].model_dump()
    missing_sku = [c.key for c in SKU_COLUMNS if c.key not in sku and c.key not in DERIVED]
    assert not missing_sku, f"Detail sheet column points at a key missing from Variant: {missing_sku}"
