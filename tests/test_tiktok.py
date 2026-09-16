"""TikTok parsers, candidate selection, export -- fixtures mirror real payloads (11/09)."""
import json
from pathlib import Path

import openpyxl
import pytest

from ecommerce.ingestion.images import thumb_source_url
from ecommerce.ingestion.raw_store import RunStore, write_json
from ecommerce.platforms.tiktok.extract.crawl import slug_allowed
from ecommerce.platforms.tiktok.extract.fetch import is_wall
from ecommerce.platforms.tiktok.parse import common as C
from ecommerce.platforms.tiktok.parse.content import description_text
from ecommerce.platforms.tiktok.parse.listing import parse_keyword_page, select_candidates
from ecommerce.platforms.tiktok.parse.product import parse_product
from ecommerce.platforms.tiktok.parse.shop import shop_address
from helpers import PROFILE, make_cfg


def parse(raw) -> dict:
    return parse_product(raw, PROFILE).to_row()

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture
def raw():
    return json.loads((FIX / "tiktok_pdp_raw.json").read_text(encoding="utf-8"))


@pytest.fixture
def components():
    return json.loads((FIX / "tiktok_keyword_components.json").read_text(encoding="utf-8"))


def test_keyword_page(components):
    page = parse_keyword_page(components)
    assert len(page["products"]) == 6
    assert page["related"] == ["giu-nhiet-om", "binh-giu-nhiet-mini", "ao-giu-nhiet"]


def test_select_candidates_filters_and_orders_by_sold(components):
    page = {"slug": "giu-nhiet", **parse_keyword_page(components)}
    again = {"slug": "binh-giu-nhiet", "products": [page["products"][0]]}
    result = select_candidates([(1, page), (2, again)], keep=10, profile=PROFILE)
    names = [c.name for c in result["kept"]]
    # shirt / lunch box dropped by type, "Bình nước thể thao" has no 'giữ nhiệt'
    assert names == ["Ly giữ nhiệt Candy 710ml có ống hút", "Bình giữ nhiệt LocknLock Metro 500ml",
                     "Cốc giữ nhiệt văn phòng 450ml"]
    reasons = {e.name: e.reason for e in result["excluded"]}
    assert "không chứa 'giữ nhiệt'" in reasons["Bình nước thể thao 1 lít"]
    assert "Quần áo" in reasons["Áo giữ nhiệt nam cổ cao"]
    assert result["kept"][1].seen_on == 2          # deduplicated across pages
    assert result["kept"][0].seller_id == "11"


def test_parse_product_core_fields(raw):
    r = parse(raw)
    assert r["item_id"] == "1730196006903450541"          # kept as text: 19 digits > Excel precision
    assert r["url"] == "https://shop.tiktok.com/vn/pdp/1730196006903450541"
    assert r["sold_total"] == 16031
    assert r["brand"] == "No brand"                       # 'KHÔNG CÓ'
    assert (r["price_min"], r["price_max"]) == (97043, 99000)
    assert r["price_before_min"] == 105000 and r["discount_pct"] == 8
    assert r["rating_star"] == 4.4 and r["rating_total"] == 1660 and r["rating_5"] == 1235
    assert r["shop_response_rate"] == 100 and r["shop_ship_48h"] == 77
    assert r["shop_location"] is None                     # 'TQ' is not an address
    assert r["origin"] == "Trung Quốc"
    assert r["warranty"] and "không" in r["warranty"].lower()
    assert r["capacity"] == "510ml"
    assert "Flash sale" in r["promotions"] and r["promotions"].count("Voucher ship") == 1
    assert r["category"] == "Thể thao & Ngoài trời > Bình nước thể thao"
    assert len(r["images"]) == 3 and r["video_count"] == 1
    assert "- Dung tích: 510ml" in r["description"] and "tos/desc1" not in r["description"]


def test_parse_product_stock_is_exact_per_sku(raw):
    r = parse(raw)
    assert r["stock_total"] == 0 + 9951 + 37
    assert r["stock_state"] == "Còn hàng (2/3 SKU còn hàng)"
    v = {x["name"]: x for x in r["variants"]}
    assert v["XANH 510"]["stock"] == 0 and v["XANH 510"]["stock_status"] == "Hết hàng"
    assert v["ĐEN 510"]["price"] == 99000 and v["ĐEN 510"]["price_before_discount"] == 120000
    assert v["XANH RÊU 510"]["tiers"] == [("PHÂN LOẠI", "XANH RÊU 510")]
    # SKU picture: sku_property_image_map first, then the option swatch
    assert "skuimg1" in v["XANH RÊU 510"]["image"] and "opt0" in v["XANH 510"]["image"]


def test_shop_address_kept_only_when_it_looks_like_one():
    attrs = {"Địa chỉ tổ chức chịu trách nhiệm hàng hóa": "Số 12 đường Lê Lợi, Quận 1, TP.HCM"}
    assert shop_address(attrs).startswith("Số 12")
    assert shop_address({"Địa chỉ tổ chức chịu trách nhiệm hàng hóa": "LocknLock"}) is None


def test_description_falls_back_to_plain_text():
    assert description_text("không phải JSON") == "không phải JSON"
    assert description_text(None) == ""


def test_missing_product_is_invalid():
    assert not C.is_valid_product({"product_info": {"error_code": 23002102, "product_info": {}}})


def test_slug_filter_uses_domain_profile_regexes():
    disc = PROFILE.discovery.tiktok
    inc, exc = disc.slug_include, disc.slug_exclude
    assert slug_allowed("ly-giu-nhiet-cam-tay", inc, exc)
    assert slug_allowed("coc-giu-nhiet-do-mixi", inc, exc)
    for bad in ("ao-giu-nhiet", "giu-nhiet-om", "hop-com-giu-nhiet", "tui-giu-nhiet", "set-do-giu-nhiet",
                "binh-sua-giu-nhiet", "binh-nuoc-nhiet"):
        assert not slug_allowed(bad, inc, exc), bad


def test_wall_detection():
    assert is_wall({"status": 200, "title": "Security Check", "components": None})
    assert is_wall({"status": 429, "title": "", "components": None})
    assert not is_wall({"status": 200, "title": "Giữ nhiệt - TikTok Shop", "components": []})
    assert not is_wall({"status": 404, "title": "404 Not Found", "components": None})


def test_tiktok_thumbnail_url_asks_cdn_for_small_jpeg():
    url = ("https://p16-oec-sg.ibyteimg.com/tos-alisg/abc~tplv-aphluv4xwc-crop-webp:1200:1200.webp"
           "?dr=15592&t=555f072d")
    small = thumb_source_url(url)
    assert "resize-jpeg:220:220.jpeg?dr=15592" in small
    assert thumb_source_url("https://down-vn.img.susercontent.com/file/x").endswith("_tn")


def test_export_end_to_end(tmp_path, raw, components):
    from ecommerce.consumption.excel.excel import write_report
    from ecommerce.platforms.tiktok.parse.dataset import build_dataset
    store = RunStore("tiktok", "t1", root=tmp_path)
    page = {"slug": "giu-nhiet", **parse_keyword_page(components)}
    write_json(store.dir / "search" / "page_001.json", page)
    cands = select_candidates([(1, page)], keep=10, profile=PROFILE)
    write_json(store.candidates_path, {"kept": [c.dump() for c in cands["kept"]],
                                       "excluded": [c.dump() for c in cands["excluded"]], "keyword": "giữ nhiệt"})
    write_json(store.item_path("7495029512148126637", "1730196006903450541"), raw)
    second = json.loads(json.dumps(raw))
    pm = second["product_info"]["product_info"]["product_model"]
    pm["product_id"], pm["sold_count"], pm["name"] = "999", "50000", "Ly giữ nhiệt Candy 710ml"
    write_json(store.item_path("1", "999"), second)
    shirt = json.loads(json.dumps(raw))
    pm = shirt["product_info"]["product_info"]["product_model"]
    pm["product_id"], pm["sold_count"], pm["name"] = "777", "900000", "Áo giữ nhiệt nam"
    write_json(store.item_path("1", "777"), shirt)

    cfg = make_cfg("tiktok", tiktok={"target": 200})
    path = write_report(build_dataset(store, cfg), cfg.export, out_dir=tmp_path, download_images=False)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["TikTok Shop", "Detail", "Checklist đề bài", "Bị loại", "Lỗi crawl", "Thông tin"]
    ws = wb["TikTok Shop"]
    headers = [c.value for c in ws[1]]
    assert "Giá sau voucher" not in headers and "Bán 30 ngày (TB tháng)" not in headers
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[headers.index("Mã SP")] for r in rows] == ["999", "1730196006903450541"]   # by sold, shirt dropped
    assert rows[0][headers.index("Hạng")] == 1
    detail = wb["Detail"]
    dh = [c.value for c in detail[1]]
    assert "Giá sau voucher" not in dh and "Tồn kho (pieces available)" in dh
    assert detail.max_row == 1 + 3 * 2
    checklist = {r[1]: r[4] for r in wb["Checklist đề bài"].iter_rows(min_row=2, values_only=True)}
    assert checklist["Tồn kho"] == "✅ Đủ"
    assert checklist["Số lượng bán trung bình"] == "❌ Không có"
    assert checklist["Tỉ lệ phản hồi"] == "✅ Đủ"


def test_tiktok_shop_link_is_not_exported(load):
    """Do not export the TikTok shop link to Excel, and do not build another one either.

    Tried 13/09: the site-provided link (shop.tiktok.com/vn/store/...) returns 404
    on desktop web; switching to www.tiktok.com/shop avoids the 404 but shows a
    blank page. On product and keyword pages NO link points to the shop -- TikTok
    Shop web has no such page. The value stays in the raw JSON; it is just kept
    out of Excel so readers do not click a broken link (decisions.md #20).
    """
    from ecommerce.consumption.excel.excel import load_layout
    from ecommerce.platforms.tiktok.parse.product import parse_product

    link = parse_product(load("tiktok_pdp_raw.json"), PROFILE).to_row()["shop_url"]
    assert link is None or link.startswith("https://shop.tiktok.com/"), link
    layout = load_layout("default")
    assert "shop_url" in layout.style("tiktok").drop_columns        # not exported to Excel
    assert "shop_url" not in layout.style("shopee").drop_columns    # Shopee's link opens fine
