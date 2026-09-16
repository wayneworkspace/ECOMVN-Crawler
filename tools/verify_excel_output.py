"""Verify an exported Excel file - the checks a DE usually runs.

Reads only the Excel file (and the raw data when present); sends NO request to the platforms.

    python tools/verify_excel_output.py output/giu_nhiet_shopee_*.xlsx
    python tools/verify_excel_output.py <file.xlsx> --sample 15   # also writes a manual check sheet

Checks:
  1. Keys and row count      - exactly 200, consecutive ranks, no duplicate ids / links
  2. Referential integrity   - every SKU on the Detail sheet belongs to a product on the main sheet
  3. Value ranges            - price > 0, stars within 0-5, stock >= 0, discount % 0-99
  4. Cross consistency       - rating total = sum of 1-5 stars; stock = sum of SKUs
  5. Column coverage         - columns that are unusually empty
  6. Outliers (IQR)          - price, SKU count, sold far from the rest
  7. Raw comparison          - re-parse the original JSON, compare cell by cell (if data/raw exists)
  8. Manual check sheet      - stratified sample by rank: 10 top, 10 middle, 10 bottom
                               (top = the products everyone looks at, middle = typical products,
                                bottom = products at the edge of the top 200, ranking bugs show here)
"""
from __future__ import annotations

import argparse
import csv
import glob
import random
import statistics
from pathlib import Path

import openpyxl

ERROR, WARNING, OK = "❌", "⚠️ ", "✅"


class Report:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def check(self, condition: bool, label: str, detail: str = "", severe: bool = True) -> None:
        if condition:
            print(f"  {OK} {label}")
            return
        (self.errors if severe else self.warnings).append(f"{label}: {detail}")
        print(f"  {ERROR if severe else WARNING} {label}" + (f" -> {detail}" if detail else ""))


def read_sheet(ws) -> tuple[list[str], list[dict]]:
    columns = [c.value for c in ws[1]]
    rows = [dict(zip(columns, [c.value for c in r], strict=False)) for r in ws.iter_rows(min_row=2)]
    return columns, rows


def outlier_bounds(values: list[float]) -> tuple[float, float]:
    """The usual IQR fence: outside [Q1 - 1.5*IQR, Q3 + 1.5*IQR] is an outlier."""
    if len(values) < 8:
        return float("-inf"), float("inf")
    q1, q3 = statistics.quantiles(sorted(values), n=4)[0], statistics.quantiles(sorted(values), n=4)[2]
    iqr = q3 - q1
    return q1 - 1.5 * iqr, q3 + 1.5 * iqr


def verify(path: Path, sample_size: int, random_sample: bool = False) -> Report:
    report = Report()
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    columns, rows = read_sheet(ws)
    detail = wb["Detail"]
    _dcols, skus = read_sheet(detail)
    print(f"\n=== {path.name}\n{len(rows)} products, {len(skus)} SKU rows, {len(columns)} columns\n")

    print("1. Keys and row count")
    ids = [r["Mã SP"] for r in rows]
    report.check(len(rows) == 200, "exactly 200 products", f"{len(rows)} rows")
    report.check([r["Hạng"] for r in rows] == list(range(1, len(rows) + 1)), "ranks consecutive 1..N")
    report.check(len(set(ids)) == len(ids), "no duplicate product id", f"{len(ids) - len(set(ids))} duplicates")
    links = [ws.cell(r, columns.index("Link sản phẩm") + 1).hyperlink for r in range(2, len(rows) + 2)]
    report.check(all(links), "every row has a clickable link")
    report.check(len({x.target for x in links if x}) == len(rows), "no duplicate link")

    print("2. Referential integrity (main sheet <-> Detail)")
    id_set = set(map(str, ids))
    unknown = {str(s["Mã SP"]) for s in skus} - id_set
    without_sku = id_set - {str(s["Mã SP"]) for s in skus}
    report.check(not unknown, "every SKU belongs to a product on the main sheet", f"{len(unknown)} unknown ids")
    report.check(not without_sku, "every product has at least one SKU", f"{len(without_sku)} products without SKU")
    sku_total = sum(r["Số SKU"] or 0 for r in rows)
    report.check(sku_total == len(skus), "sum of 'Số SKU' column = Detail row count", f"{sku_total} vs {len(skus)}")

    print("3. Value ranges")
    prices = [(r["Giá bán thấp nhất"], r["Giá bán cao nhất"]) for r in rows]
    report.check(all(a and a > 0 for a, _ in prices), "every sale price > 0")
    report.check(all(not (a and b) or a <= b for a, b in prices), "min price <= max price")
    stars = [r["Điểm đánh giá"] for r in rows if r["Điểm đánh giá"] is not None]
    report.check(all(0 <= s <= 5 for s in stars), "rating within 0-5")
    stock = [r["Tồn kho"] for r in rows if isinstance(r["Tồn kho"], (int, float))]
    report.check(all(t >= 0 for t in stock), "stock not negative")
    discounts = [r["% giảm"] for r in rows if isinstance(r["% giảm"], (int, float))]
    report.check(all(0 <= g < 100 for g in discounts), "discount % within 0-99")

    print("4. Cross consistency")
    star_mismatch = [r["Mã SP"] for r in rows
                     if r["Tổng đánh giá"]
                     and abs(r["Tổng đánh giá"] - sum(r[f"{k} sao"] or 0 for k in range(1, 6)))
                     > max(2, 0.02 * r["Tổng đánh giá"])]
    report.check(not star_mismatch, "rating total = sum of 1-5 stars", f"{len(star_mismatch)} rows differ")
    per_product: dict[str, int] = {}
    for s in skus:
        v = s.get("Tồn kho (pieces available)")
        if isinstance(v, int):
            per_product[str(s["Mã SP"])] = per_product.get(str(s["Mã SP"]), 0) + v
    stock_mismatch = [r["Mã SP"] for r in rows
                      if isinstance(r["Tồn kho"], int) and str(r["Mã SP"]) in per_product
                      and per_product[str(r["Mã SP"])] != r["Tồn kho"]]
    report.check(not stock_mismatch, "main-sheet stock = sum of SKU stock", f"{len(stock_mismatch)} rows differ")
    discount_mismatch = [r["Mã SP"] for r in rows
                         if r["Giá gốc"] and r["Giá bán thấp nhất"] and r["% giảm"]
                         and abs(round(100 * (1 - r["Giá bán thấp nhất"] / r["Giá gốc"])) - r["% giảm"]) > 1]
    report.check(not discount_mismatch, "discount % matches (original price, sale price)",
                 f"{len(discount_mismatch)} rows differ - usually the original price is the cheapest SKU's",
                 severe=False)

    print("5. Column coverage (warning when < 50%)")
    for k in columns:
        if k in ("Ảnh 1", "Ảnh 2", "Ảnh 3"):
            continue                                    # embedded pictures, the cell has no value
        filled = sum(1 for r in rows if r[k] not in (None, "", 0))
        if filled < len(rows) * 0.5:
            report.warnings.append(f"column '{k}' filled in only {filled}/{len(rows)}")
            print(f"  {WARNING} {k}: {filled}/{len(rows)} = {filled / len(rows):.0%}")

    print("6. Outliers (IQR)")
    for name, get in [("Giá bán thấp nhất", lambda r: r["Giá bán thấp nhất"]),
                      ("Số SKU", lambda r: r["Số SKU"]), ("Đã bán", lambda r: r["Đã bán"])]:
        vals = [get(r) for r in rows if isinstance(get(r), (int, float))]
        lo, hi = outlier_bounds(vals)
        far = [(r["Mã SP"], get(r), str(r["Tên sản phẩm"])[:50]) for r in rows
               if isinstance(get(r), (int, float)) and not lo <= get(r) <= hi]
        print(f"  {name}: {len(far)} products outside [{max(lo, 0):,.0f} .. {hi:,.0f}]")
        for m, v, t in far[:3]:
            print(f"      {m} | {v:,} | {t}")

    print("7. Comparison with the original JSON")
    compare_with_raw(path, rows, report)

    if sample_size:
        sheet = write_check_sheet(path, ws, columns, rows, sample_size, random_sample)
        print(f"\n8. Manual check sheet: {sheet}")
    return report


def pick_sample(rows: list[dict], per_group: int, random_sample: bool) -> list[tuple[str, dict]]:
    """Default: stratified by rank - N top, N middle, N bottom."""
    if random_sample:
        random.seed(0)
        return [("random", r) for r in random.sample(rows, min(per_group * 3, len(rows)))]
    n = len(rows)
    middle = (n - per_group) // 2
    groups = [("top", rows[:per_group]),
              ("middle", rows[middle:middle + per_group]),
              ("bottom", rows[-per_group:])]
    return [(name, r) for name, group_rows in groups for r in group_rows]


def write_check_sheet(path: Path, ws, columns: list[str], rows: list[dict], per_group: int,
                      random_sample: bool) -> Path:
    sample = pick_sample(rows, per_group, random_sample)
    sheet = path.with_name(f"manual_check_{path.stem}.csv")
    with open(sheet, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["Group", "Rank", "Product id", "Product name", "Link",
                    "Min price (file)", "Sold (file)", "Rating (file)", "SKU count (file)",
                    "Price on site", "Sold on site", "Stars on site", "Variants on site",
                    "Match? (y/n)", "Notes"])
        for group, r in sample:
            link = ws.cell(rows.index(r) + 2, columns.index("Link sản phẩm") + 1).hyperlink
            w.writerow([group, r["Hạng"], r["Mã SP"], r["Tên sản phẩm"], link.target if link else "",
                        r["Giá bán thấp nhất"], r["Đã bán"], r["Điểm đánh giá"], r["Số SKU"],
                        "", "", "", "", "", ""])
    counts = {}
    for group, _ in sample:
        counts[group] = counts.get(group, 0) + 1
    print(f"   groups: {counts} | ranks: {[r['Hạng'] for _, r in sample]}")
    return sheet


def compare_with_raw(path: Path, rows: list[dict], report: Report) -> None:
    """Re-parse the original JSON and compare cell by cell - catches bugs in the reading step."""
    try:
        from ecommerce.ingestion.raw_store import RunStore, read_json
    except ImportError:
        print("  (skipped: ecommerce is not installed)")
        return
    # giu_nhiet_<platform>_top200_<run>_<time>. The run id may carry a suffix ("20260913_2"
    # = second run of the day), so split from the RIGHT: drop the time, the rest is the run id.
    # Splitting from the left turns "20260913_2_232153" into "20260913" -> folder not found,
    # and the raw comparison is silently skipped.
    stem = path.stem
    platform = "tiktok" if "tiktok" in stem else "shopee"
    run = stem.split("_top")[1].split("_", 1)[1].rsplit("_", 1)[0] if "_top" in stem else None
    store = RunStore(platform, run) if run else RunStore.latest(platform)
    if store is None or not (store.dir / "items").is_dir():
        print(f"  (skipped: data/raw/{platform}/{run} not found)")
        return
    if platform == "tiktok":
        from ecommerce.platforms.tiktok.parse.product import parse_product
    else:
        from ecommerce.platforms.shopee.parse.product import parse_product
    in_file = {str(r["Mã SP"]): r for r in rows}
    compared_columns = {"Tên sản phẩm": "title", "Giá bán thấp nhất": "price_min", "Giá bán cao nhất": "price_max",
                        "Đã bán": "sold_total", "Điểm đánh giá": "rating_star", "Tổng đánh giá": "rating_total",
                        "Số SKU": "variant_count", "Tên shop": "shop_name"}
    compared, mismatches = 0, []
    for f in sorted((store.dir / "items").glob("*.json")):
        row = parse_product(read_json(f)).to_row()
        r = in_file.get(str(row["item_id"]))
        if r is None:
            continue
        compared += 1
        for excel_column, key in compared_columns.items():
            a, b = r[excel_column], row[key]
            if isinstance(b, float):
                b = round(b, 2)
            if isinstance(a, str) or isinstance(b, str):
                a, b = str(a).strip(), str(b).strip()
            if a != b:
                mismatches.append(f"{row['item_id']}.{excel_column}: file={str(a)[:30]} raw={str(b)[:30]}")
    if not compared:
        print("  (skipped: no product is in both the Excel file and the raw data)")
        return
    report.check(not mismatches, f"compared {compared} products with the original JSON "
                                 f"({len(compared_columns)} columns each)",
                 f"{len(mismatches)} cells differ, e.g.: {mismatches[:2]}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", help="path(s) to .xlsx files (globs allowed)")
    ap.add_argument("--sample", type=int, default=0, metavar="N",
                    help="write a manual check sheet: N products per group top / middle / bottom (e.g. --sample 10)")
    ap.add_argument("--random", action="store_true", help="random sample instead of stratified")
    args = ap.parse_args()
    total = Report()
    for pattern in args.files:
        for f in sorted(glob.glob(pattern)) or [pattern]:
            report = verify(Path(f), args.sample, args.random)
            total.errors += report.errors
            total.warnings += report.warnings
    print("\n" + "=" * 60)
    print(f"{ERROR if total.errors else OK} {len(total.errors)} errors, {len(total.warnings)} warnings")
    for x in total.errors + total.warnings:
        print("   -", x)
    raise SystemExit(1 if total.errors else 0)


if __name__ == "__main__":
    main()
